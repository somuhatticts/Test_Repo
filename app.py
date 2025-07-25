from flask import Flask, render_template, request, jsonify, session, redirect, url_for
import os
import json
from werkzeug.utils import secure_filename
import openpyxl
import xlrd
import csv
import io
import tempfile
from onedrive_service import OneDriveService, download_public_onedrive_file
from config import SECRET_KEY

app = Flask(__name__)
app.config['UPLOAD_FOLDER'] = 'uploads'
app.config['MAX_CONTENT_LENGTH'] = 16 * 1024 * 1024  # 16MB max file size
app.secret_key = SECRET_KEY

# Create uploads directory if it doesn't exist
os.makedirs(app.config['UPLOAD_FOLDER'], exist_ok=True)

ALLOWED_EXTENSIONS = {'xlsx', 'xls', 'csv'}

def allowed_file(filename):
    return '.' in filename and filename.rsplit('.', 1)[1].lower() in ALLOWED_EXTENSIONS

def detect_cell_type(value):
    """Detect the type of a cell value"""
    if value is None or value == '':
        return 'empty'
    elif isinstance(value, (int, float)):
        return 'number'
    elif isinstance(value, str):
        # Try to convert to number
        try:
            float(value)
            return 'number'
        except (ValueError, TypeError):
            return 'text'
    else:
        return 'other'

def process_xlsx_file(filepath):
    """Process .xlsx file using openpyxl"""
    workbook = openpyxl.load_workbook(filepath)
    sheet = workbook.active
    
    # Get data dimensions
    max_row = min(sheet.max_row, 100)  # Limit to 100 rows
    max_col = min(sheet.max_column, 20)  # Limit to 20 columns
    
    # Extract column headers
    columns = []
    for col in range(1, max_col + 1):
        cell_value = sheet.cell(row=1, column=col).value
        columns.append(str(cell_value) if cell_value is not None else f'Column {col}')
    
    # Extract data
    data = []
    for row in range(2, max_row + 1):  # Start from row 2 (skip header)
        row_data = []
        for col in range(1, max_col + 1):
            cell_value = sheet.cell(row=row, column=col).value
            cell_type = detect_cell_type(cell_value)
            row_data.append({
                'value': cell_value if cell_value is not None else '',
                'type': cell_type
            })
        data.append(row_data)
    
    # Calculate statistics
    total_rows = max_row - 1  # Exclude header
    total_cols = max_col
    numeric_cols = 0
    text_cols = 0
    
    # Count column types based on first data row
    if data:
        for cell in data[0]:
            if cell['type'] == 'number':
                numeric_cols += 1
            elif cell['type'] == 'text':
                text_cols += 1
    
    return {
        'data': data,
        'columns': columns,
        'info': {
            'rows': total_rows,
            'columns': total_cols,
            'numeric_columns': numeric_cols,
            'text_columns': text_cols
        }
    }

def process_xls_file(filepath):
    """Process .xls file using xlrd"""
    workbook = xlrd.open_workbook(filepath)
    sheet = workbook.sheet_by_index(0)
    
    # Get data dimensions
    max_row = min(sheet.nrows, 100)  # Limit to 100 rows
    max_col = min(sheet.ncols, 20)  # Limit to 20 columns
    
    # Extract column headers
    columns = []
    if max_row > 0:
        for col in range(max_col):
            cell_value = sheet.cell_value(0, col)
            columns.append(str(cell_value) if cell_value else f'Column {col + 1}')
    
    # Extract data
    data = []
    for row in range(1, max_row):  # Start from row 1 (skip header)
        row_data = []
        for col in range(max_col):
            cell_value = sheet.cell_value(row, col)
            cell_type = detect_cell_type(cell_value)
            row_data.append({
                'value': cell_value if cell_value != '' else '',
                'type': cell_type
            })
        data.append(row_data)
    
    # Calculate statistics
    total_rows = max_row - 1 if max_row > 0 else 0
    total_cols = max_col
    numeric_cols = 0
    text_cols = 0
    
    # Count column types based on first data row
    if data:
        for cell in data[0]:
            if cell['type'] == 'number':
                numeric_cols += 1
            elif cell['type'] == 'text':
                text_cols += 1
    
    return {
        'data': data,
        'columns': columns,
        'info': {
            'rows': total_rows,
            'columns': total_cols,
            'numeric_columns': numeric_cols,
            'text_columns': text_cols
        }
    }

def process_csv_content(csv_content):
    """Process CSV content from bytes or string"""
    data = []
    columns = []
    
    # Convert bytes to string if necessary
    if isinstance(csv_content, bytes):
        csv_content = csv_content.decode('utf-8')
    
    # Create a StringIO object from the content
    csv_file = io.StringIO(csv_content)
    
    # Try to detect delimiter
    sample = csv_content[:1024]
    sniffer = csv.Sniffer()
    try:
        delimiter = sniffer.sniff(sample).delimiter
    except:
        delimiter = ','  # Default to comma
    
    reader = csv.reader(csv_file, delimiter=delimiter)
    
    # Read header
    try:
        columns = next(reader)
        columns = [str(col).strip() for col in columns[:20]]  # Limit to 20 columns
    except StopIteration:
        columns = []
    
    # Read data rows
    row_count = 0
    for row in reader:
        if row_count >= 99:  # Limit to 99 data rows + 1 header = 100 total
            break
        
        row_data = []
        for i, cell_value in enumerate(row[:20]):  # Limit to 20 columns
            cell_value = str(cell_value).strip() if cell_value else ''
            cell_type = detect_cell_type(cell_value)
            row_data.append({
                'value': cell_value,
                'type': cell_type
            })
        
        # Pad row if it has fewer columns than header
        while len(row_data) < len(columns):
            row_data.append({'value': '', 'type': 'empty'})
        
        data.append(row_data)
        row_count += 1
    
    # Calculate statistics
    total_rows = len(data)
    total_cols = len(columns)
    numeric_cols = 0
    text_cols = 0
    
    # Count column types based on first data row
    if data:
        for cell in data[0]:
            if cell['type'] == 'number':
                numeric_cols += 1
            elif cell['type'] == 'text':
                text_cols += 1
    
    return {
        'data': data,
        'columns': columns,
        'info': {
            'rows': total_rows,
            'columns': total_cols,
            'numeric_columns': numeric_cols,
            'text_columns': text_cols
        }
    }

def process_csv_file(filepath):
    """Process .csv file"""
    with open(filepath, 'r', encoding='utf-8', newline='') as file:
        content = file.read()
    return process_csv_content(content)

def get_file_type(filename):
    """Get file type based on extension"""
    ext = filename.lower().split('.')[-1]
    if ext == 'csv':
        return 'csv'
    elif ext in ['xlsx', 'xls']:
        return 'excel'
    else:
        return 'unknown'

@app.route('/')
def index():
    return render_template('index.html')

@app.route('/auth/login')
def login():
    """Initiate OneDrive authentication"""
    try:
        onedrive_service = OneDriveService()
        auth_url = onedrive_service.get_auth_url()
        return redirect(auth_url)
    except Exception as e:
        return jsonify({'error': f'Authentication setup failed: {str(e)}'}), 500

@app.route('/auth/callback')
def auth_callback():
    """Handle OAuth callback from Microsoft"""
    try:
        auth_code = request.args.get('code')
        if not auth_code:
            return jsonify({'error': 'No authorization code received'}), 400
        
        onedrive_service = OneDriveService()
        if onedrive_service.get_token_from_code(auth_code):
            session['access_token'] = onedrive_service.access_token
            return redirect(url_for('index'))
        else:
            return jsonify({'error': 'Failed to get access token'}), 400
    except Exception as e:
        return jsonify({'error': f'Authentication failed: {str(e)}'}), 500

@app.route('/onedrive/files')
def list_onedrive_files():
    """List files from OneDrive"""
    try:
        if 'access_token' not in session:
            return jsonify({'error': 'Not authenticated. Please login first.'}), 401
        
        onedrive_service = OneDriveService()
        onedrive_service.access_token = session['access_token']
        
        # Search for Excel and CSV files
        files = onedrive_service.search_files("", ['.xlsx', '.xls', '.csv'])
        
        # Format files for frontend
        formatted_files = []
        for file in files:
            if 'file' in file:  # Only include actual files, not folders
                formatted_files.append({
                    'id': file['id'],
                    'name': file['name'],
                    'size': file.get('size', 0),
                    'modified': file.get('lastModifiedDateTime', ''),
                    'downloadUrl': file.get('@microsoft.graph.downloadUrl', ''),
                    'type': get_file_type(file['name'])
                })
        
        return jsonify({'files': formatted_files})
    except Exception as e:
        return jsonify({'error': f'Failed to list files: {str(e)}'}), 500

@app.route('/onedrive/browse')
def browse_onedrive():
    """Browse OneDrive folders and files"""
    try:
        if 'access_token' not in session:
            return jsonify({'error': 'Not authenticated. Please login first.'}), 401
        
        folder_path = request.args.get('path', '')
        file_type = request.args.get('type', 'all')  # 'csv', 'excel', or 'all'
        
        onedrive_service = OneDriveService()
        onedrive_service.access_token = session['access_token']
        
        # Get folders for navigation
        folders = onedrive_service.get_folders(folder_path)
        
        # Get files based on type filter
        if file_type == 'csv':
            files = onedrive_service.get_csv_files(folder_path)
        elif file_type == 'excel':
            files = onedrive_service.search_files("", ['.xlsx', '.xls'], folder_path)
        else:
            files = onedrive_service.search_files("", ['.xlsx', '.xls', '.csv'], folder_path)
        
        # Format files for frontend
        formatted_files = []
        for file in files:
            if 'file' in file:
                formatted_files.append({
                    'id': file['id'],
                    'name': file['name'],
                    'size': file.get('size', 0),
                    'modified': file.get('lastModifiedDateTime', ''),
                    'type': get_file_type(file['name']),
                    'path': file.get('parentReference', {}).get('path', '') + '/' + file['name']
                })
        
        # Format folders for frontend
        formatted_folders = []
        for folder in folders:
            formatted_folders.append({
                'id': folder['id'],
                'name': folder['name'],
                'path': folder['path'],
                'childCount': folder['childCount'],
                'type': 'folder'
            })
        
        return jsonify({
            'currentPath': folder_path,
            'folders': formatted_folders,
            'files': formatted_files,
            'totalFiles': len(formatted_files),
            'totalFolders': len(formatted_folders)
        })
    except Exception as e:
        return jsonify({'error': f'Failed to browse OneDrive: {str(e)}'}), 500

@app.route('/onedrive/csv-files')
def list_csv_files():
    """List specifically CSV files from OneDrive"""
    try:
        if 'access_token' not in session:
            return jsonify({'error': 'Not authenticated. Please login first.'}), 401
        
        folder_path = request.args.get('path', '')
        
        onedrive_service = OneDriveService()
        onedrive_service.access_token = session['access_token']
        
        # Get CSV files
        csv_files = onedrive_service.get_csv_files(folder_path)
        
        # Format files for frontend
        formatted_files = []
        for file in csv_files:
            formatted_files.append({
                'id': file['id'],
                'name': file['name'],
                'size': file.get('size', 0),
                'modified': file.get('lastModifiedDateTime', ''),
                'type': 'csv',
                'path': file.get('parentReference', {}).get('path', '') + '/' + file['name']
            })
        
        return jsonify({
            'files': formatted_files,
            'count': len(formatted_files),
            'currentPath': folder_path
        })
    except Exception as e:
        return jsonify({'error': f'Failed to list CSV files: {str(e)}'}), 500

@app.route('/onedrive/download/<file_id>')
def download_onedrive_file(file_id):
    """Download and process a file from OneDrive"""
    try:
        if 'access_token' not in session:
            return jsonify({'error': 'Not authenticated. Please login first.'}), 401
        
        onedrive_service = OneDriveService()
        onedrive_service.access_token = session['access_token']
        
        # Download file content
        file_content = onedrive_service.download_file(file_id)
        
        # Create temporary file to process
        with tempfile.NamedTemporaryFile(delete=False) as temp_file:
            temp_file.write(file_content)
            temp_filepath = temp_file.name
        
        try:
            # Determine file type and process accordingly
            # For simplicity, we'll assume it's CSV if it contains text-like content
            try:
                # Try to decode as text (CSV)
                content_str = file_content.decode('utf-8')
                processed_data = process_csv_content(content_str)
            except UnicodeDecodeError:
                # If it's not text, try as Excel file
                if file_content.startswith(b'PK'):  # ZIP signature (xlsx files)
                    processed_data = process_xlsx_file(temp_filepath)
                else:
                    # Try as xls file
                    processed_data = process_xls_file(temp_filepath)
            
            return jsonify({
                'success': True,
                'data': processed_data
            })
        finally:
            # Clean up temporary file
            os.unlink(temp_filepath)
            
    except Exception as e:
        return jsonify({'error': f'Failed to download file: {str(e)}'}), 500

@app.route('/onedrive/demo')
def demo_onedrive():
    """Demo endpoint that uses sample data (for testing without full OAuth setup)"""
    try:
        # Use the demo function to get sample CSV content
        sample_content = download_public_onedrive_file("demo://sample.csv")
        processed_data = process_csv_content(sample_content)
        
        return jsonify({
            'success': True,
            'data': processed_data,
            'message': 'Demo data from OneDrive simulation'
        })
    except Exception as e:
        return jsonify({'error': f'Demo failed: {str(e)}'}), 500

@app.route('/upload', methods=['POST'])
def upload_file():
    """Original file upload functionality (still available)"""
    if 'file' not in request.files:
        return jsonify({'error': 'No file selected'}), 400
    
    file = request.files['file']
    if file.filename == '':
        return jsonify({'error': 'No file selected'}), 400
    
    if file and allowed_file(file.filename):
        try:
            filename = secure_filename(file.filename)
            filepath = os.path.join(app.config['UPLOAD_FOLDER'], filename)
            file.save(filepath)
            
            # Process the file based on extension
            file_ext = filename.rsplit('.', 1)[1].lower()
            
            if file_ext == 'xlsx':
                processed_data = process_xlsx_file(filepath)
            elif file_ext == 'xls':
                processed_data = process_xls_file(filepath)
            elif file_ext == 'csv':
                processed_data = process_csv_file(filepath)
            else:
                return jsonify({'error': 'Unsupported file format'}), 400
            
            # Clean up uploaded file
            os.remove(filepath)
            
            return jsonify({
                'success': True,
                'data': processed_data
            })
            
        except Exception as e:
            return jsonify({'error': f'Error processing file: {str(e)}'}), 500
    
    return jsonify({'error': 'Invalid file type. Please upload Excel (.xlsx, .xls) or CSV files.'}), 400

if __name__ == '__main__':
    app.run(debug=True, host='0.0.0.0', port=5000)